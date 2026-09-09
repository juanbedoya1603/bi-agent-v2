from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMN_LABELS = {
    "sales": "Ventas",
    "units": "Unidades",
    "previous_sales": "Ventas - período anterior",
    "current_sales": "Ventas - período actual",
    "previous_units": "Unidades - período anterior",
    "current_units": "Unidades - período actual",
    "growth": "Variación",
    "mom_growth": "Variación mensual",
    "productName": "Producto",
    "brandName": "Marca",
    "categoryName": "Categoría",
    "cityName": "Ciudad",
    "tickets": "Tickets",
    "stores": "Tiendas",
}


def readable_header(column: str) -> str:
    if column in COLUMN_LABELS:
        return COLUMN_LABELS[column]
    return column.replace("_", " ").strip().capitalize()


def _safe_cell(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def build_excel(columns: list[str], rows: list[list[Any]]) -> bytes:
    """Crea una sola hoja en memoria a partir de las filas ya visibles."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Datos"
    sheet.freeze_panes = "A2"

    headers = [readable_header(column) for column in columns]
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_cell(value) for value in row])

    header_fill = PatternFill("solid", fgColor="0F5C5A")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.auto_filter.ref = sheet.dimensions
    for column_index, header in enumerate(headers, start=1):
        values = [
            header,
            *(
                "" if row[column_index - 1] is None else str(row[column_index - 1])
                for row in rows
            ),
        ]
        width = min(max(len(value) for value in values) + 2, 48)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    # Verifica que el stream producido sea un XLSX legible antes de responderlo.
    load_workbook(output, read_only=True).close()
    return output.getvalue()
